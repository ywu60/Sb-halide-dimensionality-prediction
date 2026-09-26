"""Export verified extraction records as final Excel and JSON datasets.

This is intentionally a one-step export. Human review is optional and happens
by opening and editing the resulting workbook; the pipeline does not require
review statuses or a separate apply-review command.

Writes data/output/final_dataset.xlsx and data/output/final_dataset.json.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from src.utils import get_logger, load_config, read_jsonl

logger = get_logger("export_final_dataset")

FINAL_COLUMNS: list[tuple[str, str, int, bool]] = [
    ("paper_id", "paper_id", 10, False),
    ("compound_id", "compound_id", 14, False),
    ("label_in_paper", "label_in_paper", 24, True),
    ("compound_name_reported", "compound_name_reported", 32, True),
    ("compound_formula_reported", "compound_formula_reported", 26, False),
    ("formula_abbreviation_glossary", "formula_abbreviation_glossary", 38, True),
    ("cation_name_reported", "cation_name_reported", 24, True),
    ("cation_abbreviation_reported", "cation_abbreviation_reported", 18, False),
    ("cation_formula_explicit", "cation_formula_explicit", 22, False),
    ("cation_formula_status", "cation_formula_status", 16, False),
    ("cation_evidence_text", "cation_evidence_text", 55, True),
    ("cation_source", "cation_source", 26, True),
    ("sb_oxidation_state_reported", "sb_oxidation_state_reported", 16, False),
    ("halides_present", "halides_present", 14, False),
    ("halides_bonded_to_sb", "halides_bonded_to_sb", 18, False),
    ("sb_halide_unit_reported", "sb_halide_unit_reported", 32, True),
    ("sb_halide_connectivity_evidence_text", "sb_halide_connectivity_evidence_text", 55, True),
    ("connectivity_source", "connectivity_source", 26, True),
    ("sb_halide_dimensionality_llm", "sb_halide_dimensionality", 18, False),
    ("dimensionality_reasoning", "dimensionality_reasoning", 55, True),
    ("dimensionality_evidence_status", "dimensionality_evidence_status", 18, False),
    ("dimensionality_review_reason", "dimensionality_review_reason", 30, True),
    ("synthesis_target_compound", "synthesis_target_compound", 24, True),
    ("synthesis_evidence_text", "synthesis_evidence_text", 55, True),
    ("synthesis_source", "synthesis_source", 26, True),
    ("verification_status", "verification_status", 18, False),
    ("verification_notes", "verification_notes", 40, True),
    ("automatic_flags", "automatic_flags", 36, True),
    ("retrieval_cycles_used", "retrieval_cycles_used", 14, False),
]


def _joinlist(value):
    if isinstance(value, list):
        return "; ".join(str(item) for item in value if item is not None)
    return value


def build_final_records(records: list[dict]) -> list[dict]:
    selected = [
        {output: record.get(source) for source, output, _, _ in FINAL_COLUMNS}
        for record in records
    ]
    return sorted(selected, key=lambda record: (record["paper_id"], record["compound_id"]))


def build_final_dataframe(records: list[dict]) -> pd.DataFrame:
    df = pd.DataFrame(build_final_records(records))
    for column in df.columns:
        df[column] = df[column].apply(_joinlist)
    return df


def main() -> None:
    cfg = load_config()
    data_dir = Path(cfg["paths"]["data_dir"])
    records = list(read_jsonl(data_dir / "records" / "08_verified_records.jsonl"))
    if not records:
        raise SystemExit("data/records/08_verified_records.jsonl is empty — run verify first")

    final_records = build_final_records(records)
    final_df = build_final_dataframe(records)
    output_dir = data_dir / "output"
    excel_path = output_dir / "final_dataset.xlsx"
    json_path = output_dir / "final_dataset.json"
    output_dir.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        final_df.to_excel(writer, sheet_name="Final Dataset", index=False)
        ws = writer.sheets["Final Dataset"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for index, (_, _, width, wrap) in enumerate(FINAL_COLUMNS, start=1):
            letter = get_column_letter(index)
            ws.column_dimensions[letter].width = width
            ws[f"{letter}1"].font = Font(bold=True)
            if wrap:
                for row in range(2, ws.max_row + 1):
                    ws[f"{letter}{row}"].alignment = Alignment(wrap_text=True, vertical="top")

    json_path.write_text(
        json.dumps(final_records, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    logger.info(
        "wrote final dataset with %d compounds -> %s and %s",
        len(final_df),
        excel_path,
        json_path,
    )


if __name__ == "__main__":
    main()
